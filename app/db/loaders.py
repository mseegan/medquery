"""Readers for the source records: the patient spreadsheet and visit-note PDFs."""
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

SECTION_RE = re.compile(r"^(Subjective|Objective|Assessment|Plan) Notes:?[ \t]*$", re.M)
FOOTER_RE = re.compile(r"^(Patient MRN:.*|Encounter ID:.*|Page \d+ of \d+)\s*$", re.M)
ENCOUNTER_HEADER_RE = re.compile(
    r"^(?P<name>[A-Za-z][A-Za-z .'-]+?)[ \t]*\n"
    r"(?P<when>\d{1,2}/\d{1,2}/\d{4} \d{1,2}:\d{2} [AP]M)[ \t]*\n"
    r"Location:[ \t]*(?P<location>.*?)[ \t]*\n"
    r"Patient #:[ \t]*(?P<mrn>\S+)[ \t]*\n"
    r"DOB:[ \t]*(?P<dob>\S+)[ \t]*\n"
    r"Gender:[ \t]*(?P<gender>\w+)",
    re.M,
)


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _us_date(value: str) -> date:
    return datetime.strptime(value, "%m/%d/%Y").date()


def load_patients_xlsx(path: Path) -> list[dict]:
    """Read patient rows, merging repeated rows for the same person."""
    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(h).strip().lower() for h in rows[0]]

    merged: dict[str, dict] = {}
    for row in rows[1:]:
        record = {
            key: (str(value).strip() if value is not None else "")
            for key, value in zip(header, row)
        }
        name = record.get("name", "")
        if not name:
            continue
        entry = merged.setdefault(name.lower(), {})
        for key, value in record.items():
            if value and not entry.get(key):
                entry[key] = value

    return [
        {
            "name": e["name"],
            "phone": e.get("phone_number", ""),
            "email": e.get("email", ""),
            "age": int(float(e["age"])) if e.get("age") else None,
            "gender": e.get("gender", ""),
            "address": e.get("address", ""),
            "summary": e.get("summary", ""),
        }
        for e in merged.values()
    ]


def _split_sections(block: str) -> dict[str, str]:
    parts = SECTION_RE.split(block)
    return {parts[i].lower(): _clean(parts[i + 1]) for i in range(1, len(parts) - 1, 2)}


def _diagnoses(assessment: str) -> tuple[str, str]:
    """Diagnosis names and ICD-10 codes, joined with '; '."""
    coded = re.findall(r"([^\[\],]+?)\s*\[ICD-10:\s*([A-Z0-9.]+)\]", assessment)
    if coded:
        names = [re.sub(r"^DIAGNOSIS:\s*", "", n, flags=re.I).strip() for n, _ in coded]
        return "; ".join(names), "; ".join(code for _, code in coded)
    match = re.search(r"Diagnosis:\s*(.+?)\s*\(([A-Z]\d{2}[\w.]*)\)", assessment)
    if match:
        return match.group(1).strip(), match.group(2)
    return re.sub(r"^Diagnosis:\s*", "", assessment, flags=re.I).strip(), ""


def _encounter(visit_date: date, sections: dict[str, str], provider: str, complaint: str) -> dict:
    diagnosis, icd10 = _diagnoses(sections.get("assessment", ""))
    return {
        "visit_date": visit_date,
        "provider": provider,
        "chief_complaint": complaint,
        "subjective": sections.get("subjective", ""),
        "objective": sections.get("objective", ""),
        "assessment": sections.get("assessment", ""),
        "plan": sections.get("plan", ""),
        "diagnosis": diagnosis,
        "icd10_code": icd10,
    }


def _parse_simple_note(text: str) -> tuple[dict, list[dict]]:
    """A one-visit note with 'Patient:', 'DOB:', and 'Visit Date:' header lines."""
    def field(label: str) -> str:
        match = re.search(rf"^{label}:[ \t]*(.+)$", text, re.M)
        return match.group(1).strip() if match else ""

    patient = {
        "name": field("Patient"),
        "dob": _us_date(field("DOB")) if field("DOB") else None,
        "gender": field("Gender"),
        "phone": field("Phone"),
        "address": field("Address"),
        "mrn": "",
    }
    visit = _us_date(field("Visit Date"))
    return patient, [_encounter(visit, _split_sections(text), "", "")]


def _parse_chart_note(text: str) -> tuple[dict, list[dict]]:
    """A chart export holding one or more visits, each with a name/date/location
    header, SOAP sections, and a trailing 'Encounter' block."""
    text = FOOTER_RE.sub("", text)
    headers = list(ENCOUNTER_HEADER_RE.finditer(text))

    patient: dict = {}
    encounters = []
    for i, header in enumerate(headers):
        end = headers[i + 1].start() if i + 1 < len(headers) else len(text)
        block = text[header.end() : end]
        if "Subjective Notes" not in block:
            continue  # history/lab pages repeat the header without a visit note

        soap, *rest = re.split(r"\nEncounter[ \t]*\n", block, maxsplit=1)
        trailer = rest[0] if rest else ""
        provider = re.search(r"Signed electronically by ([^,(]+)", trailer)
        complaint = re.search(r"ChiefComplaint:\s*(.+?)\s+ACTIVE", _clean(trailer))
        complaint_text = re.sub(r"(-zocdoc)?\s*-\s*[A-Z]{2}$", "", complaint.group(1)) if complaint else ""

        patient = patient or {
            "name": header["name"].strip(),
            "dob": _us_date(header["dob"]),
            "gender": header["gender"],
            "phone": "",
            "address": "",
            "mrn": header["mrn"],
        }
        encounters.append(
            _encounter(
                datetime.strptime(header["when"], "%m/%d/%Y %I:%M %p").date(),
                _split_sections(soap),
                provider.group(1).strip() if provider else "",
                complaint_text.strip(),
            )
        )
    return patient, encounters


def parse_note_pdf(path: Path) -> tuple[dict, list[dict]]:
    """Read a visit-note PDF into (patient details, encounters)."""
    text = "\n".join(page.extract_text() for page in PdfReader(path).pages)
    if "Location:" in text and "Patient #:" in text:
        return _parse_chart_note(text)
    return _parse_simple_note(text)
